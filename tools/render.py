"""Rendu visuel : PNG par slide + planches de contrôle.

Moteur : PowerPoint (COM, Windows) ; à défaut LibreOffice (soffice -> PDF -> PNG via pdftoppm ou
PyMuPDF), ex. dans le skill claude.ai. Le rendu LibreOffice est approximatif (police Aptos remplacée).

  python tools/render.py png "<pptx|potx>" [--slides 3,5-8]   -> work/<doc>/png/<nom>/NNN.png + planche
  python tools/render.py compare "<rendu 04 .pptx>"             -> work/<doc>/compare/NNN.png
        (source ↔ rendu côte à côte, d'après build_log.json)
  python tools/render.py layouts                                -> tools/catalog/layouts/*.png

Les slides masquées sont aussi exportées (marquées « MASQUÉE »).
"""
import argparse
import json
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from common import (CORRECTED, REF_EN, RENDER, ROOT, SRC_FR, TEMPLATE, WORK, as_pptx, open_prs,
                    work_dir)

W, H = 1280, 720


def _font(size=22):
    for f in ("arialbd.ttf", "arial.ttf"):
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def parse_range(txt, n):
    if not txt:
        return list(range(1, n + 1))
    out = []
    for part in txt.split(","):
        if "-" in part:
            a, b = part.split("-")
            out += range(int(a), int(b) + 1)
        else:
            out.append(int(part))
    return [i for i in out if 1 <= i <= n]


def _export_png_libreoffice(pptx: Path, out_dir: Path, slides=None):
    """Repli sans PowerPoint : LibreOffice -> PDF -> PNG. Les slides masquées sont démasquées dans une
    copie pour que chaque page du PDF corresponde à une slide."""
    from pptx import Presentation
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("Rendu impossible : ni PowerPoint ni LibreOffice (soffice) disponibles")
    prs = Presentation(str(pptx))
    for s in prs.slides:
        s._element.attrib.pop("show", None)
    res = {}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        copy = tmp / "render.pptx"
        prs.save(str(copy))
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(copy)],
                       check=True, capture_output=True, timeout=900)
        pdf = tmp / "render.pdf"
        if shutil.which("pdftoppm"):
            subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(W), "-scale-to-y", str(H),
                            str(pdf), str(tmp / "p")], check=True, capture_output=True, timeout=900)
            pages = sorted(tmp.glob("p-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
        else:
            import fitz  # PyMuPDF
            pages = []
            for i, page in enumerate(fitz.open(str(pdf)), 1):
                pix = page.get_pixmap(matrix=fitz.Matrix(W / page.rect.width, H / page.rect.height))
                pages.append(tmp / f"p-{i}.png")
                pix.save(str(pages[-1]))
        for i, pg in enumerate(pages, 1):
            if slides and i not in slides:
                continue
            target = out_dir / f"{i:03d}.png"
            shutil.move(str(pg), str(target))
            res[i] = target
    return res


def export_png(pptx: Path, out_dir: Path, slides=None, force=False):
    """Exporte les slides en PNG (PowerPoint, sinon LibreOffice). Renvoie {n: chemin}."""
    pptx = as_pptx(pptx).resolve()
    out_dir = Path(out_dir).resolve()          # PowerPoint exige un chemin absolu
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return _export_png_libreoffice(pptx, out_dir, slides)
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    res = {}
    try:
        pres = app.Presentations.Open(str(pptx), ReadOnly=True, Untitled=True, WithWindow=False)
        try:
            n = pres.Slides.Count
            for i in (slides or range(1, n + 1)):
                target = out_dir / f"{i:03d}.png"
                if force or not target.exists() or target.stat().st_mtime < pptx.stat().st_mtime:
                    pres.Slides(i).Export(str(target), "PNG", W, H)
                res[i] = target
        finally:
            pres.Close()
    finally:
        try:
            if app.Presentations.Count == 0:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return res


def label(img, text, color=(186, 19, 26)):
    d = ImageDraw.Draw(img)
    f = _font(22)
    d.rectangle([0, 0, d.textlength(text, font=f) + 20, 34], fill=color)
    d.text((10, 5), text, fill="white", font=f)
    return img


def contact_sheet(pngs: dict, hidden: set, out: Path, cols=5):
    thumbs = []
    for n, p in sorted(pngs.items()):
        im = Image.open(p).convert("RGB").resize((384, 216))
        label(im, f"{n}" + (" · MASQUÉE" if n in hidden else ""))
        thumbs.append(im)
    rows = (len(thumbs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * 394 + 10, rows * 226 + 10), "white")
    for k, im in enumerate(thumbs):
        sheet.paste(im, (10 + (k % cols) * 394, 10 + (k // cols) * 226))
    sheet.save(out)
    return out


def cmd_png(path: Path, slides_txt=None):
    prs = open_prs(path)
    n = len(prs.slides)
    hidden = {i for i, s in enumerate(prs.slides, 1) if s._element.get("show") == "0"}
    root = next((n for b, n in ((SRC_FR, "02"), (REF_EN, "03"), (RENDER, "04"), (CORRECTED, "05"))
                 if b.resolve() in path.resolve().parents), "autre")
    d = work_dir(path) / "png" / f"{root} - {path.stem}"
    pngs = export_png(path, d, parse_range(slides_txt, n))
    sheet = contact_sheet(pngs, hidden, d / "_planche.png")
    print(f"{len(pngs)} PNG -> {d}\nPlanche : {sheet}")


def cmd_compare(render: Path):
    render = render.resolve()
    logs = [p for p in WORK.rglob("build_log.json")
            if Path(json.loads(p.read_text(encoding="utf-8"))["output"]).resolve() == render]
    if not logs:
        sys.exit("build_log.json introuvable pour ce rendu")
    log = json.loads(logs[0].read_text(encoding="utf-8"))
    source = Path(log["source"])
    wd = logs[0].parent
    out_png = export_png(render, wd / "png" / render.stem)
    src_png = export_png(source, wd / "png" / "_source")
    cmp_dir = wd / "compare"
    cmp_dir.mkdir(exist_ok=True)
    for e in log["slides"]:
        o = Image.open(out_png[e["out"]]).convert("RGB").resize((960, 540))
        label(o, f"RENDU {e['out']} · {e['mode']} · {e.get('layout') or ''}")
        if e.get("src"):
            s = Image.open(src_png[e["src"]]).convert("RGB").resize((960, 540))
            label(s, f"SOURCE {e['src']}", color=(60, 60, 60))
        else:
            s = Image.new("RGB", (960, 540), (240, 240, 240))
            label(s, "(slide générée)", color=(120, 120, 120))
        sheet = Image.new("RGB", (1940, 560), "white")
        sheet.paste(s, (10, 10)); sheet.paste(o, (970, 10))
        sheet.save(cmp_dir / f"{e['out']:03d}.png")
    print(f"{len(log['slides'])} comparaisons -> {cmp_dir}")


def cmd_layouts():
    """Masques annotés (zones idx) + slides de design du template -> tools/catalog/."""
    from pptxops import delete_slide
    prs = open_prs(TEMPLATE)
    tpl = list(prs.slides)
    for lay in prs.slide_layouts:
        s = prs.slides.add_slide(lay)
        for ph in s.placeholders:
            if ph.has_text_frame:
                ph.text_frame.text = f"idx {ph.placeholder_format.idx}"
        if s.shapes.title is not None:
            s.shapes.title.text_frame.text = lay.name
    for s in tpl:
        delete_slide(prs, s)
    tmp = WORK / "_cache" / "layouts_demo.pptx"
    tmp.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(tmp))
    cat = Path(__file__).parent / "catalog"
    d = cat / "layouts"
    pngs = export_png(tmp, d, force=True)
    contact_sheet(pngs, set(), d / "_planche.png", cols=4)
    print(f"{len(pngs)} masques rendus -> {d}")
    dd = cat / "designs"
    dpngs = export_png(TEMPLATE, dd, force=True)
    contact_sheet(dpngs, set(), dd / "_planche.png", cols=5)
    print(f"{len(dpngs)} slides de design rendues -> {dd}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["png", "compare", "layouts"])
    ap.add_argument("path", nargs="?")
    ap.add_argument("--slides")
    a = ap.parse_args()
    if a.cmd == "png":
        cmd_png(Path(a.path).resolve(), a.slides)
    elif a.cmd == "compare":
        cmd_compare(Path(a.path))
    else:
        cmd_layouts()


if __name__ == "__main__":
    main()
