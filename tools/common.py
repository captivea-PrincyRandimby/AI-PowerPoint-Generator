"""Constantes et utilitaires partagés par tous les scripts du projet."""
from __future__ import annotations

import copy
import hashlib
import os
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

# Console Windows : forcer l'UTF-8 pour les accents
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

# --------------------------------------------------------------------------
# Chemins
# --------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parent.parent
PROMPTS = ROOT / "00. Prompts"
TEMPLATE_DIR = ROOT / "01. Template"
TEMPLATE_NAME = "2026-template-captivea.potx"   # template de référence (dernière version)


def _find_template() -> Path:
    """Le template de référence ; à défaut, le seul .potx/.pptx présent dans 01. Template."""
    ref = TEMPLATE_DIR / TEMPLATE_NAME
    if ref.exists():
        return ref
    cands = sorted((p for p in TEMPLATE_DIR.glob("*") if p.suffix.lower() in (".potx", ".pptx")
                    and not p.name.startswith("~$")), key=lambda p: p.stat().st_mtime, reverse=True)
    return cands[0] if cands else ref


TEMPLATE = _find_template()
SRC_FR = ROOT / "02. French PPT to transform"
REF_EN = ROOT / "03. English PPT for translation"
RENDER = ROOT / "04. Render from claude"
CORRECTED = ROOT / "05. Render corrected"
UNIQUE_FR = RENDER / "unique_slides"
UNIQUE_EN = RENDER / "unique_slides_EN"
TRACK_FR = RENDER / "Slides uniques - suivi.xlsx"
TRACK_EN = RENDER / "Slides uniques EN - suivi.xlsx"
# inventaires, plans, rendus PNG (hors livrables) ; CAPTIVEA_WORK le déplace (skill claude.ai : dossier
# du skill en lecture seule)
WORK = Path(os.environ["CAPTIVEA_WORK"]) if os.environ.get("CAPTIVEA_WORK") else ROOT / "work"
CACHE = WORK / "_cache"       # .potx convertis en .pptx

PPT_EXT = {".pptx", ".potx"}

# --------------------------------------------------------------------------
# Charte 2026
# --------------------------------------------------------------------------
RED = "BA131A"
RED2 = "CD5358"
GRAD_END = "E2575D"
CHARTE_COLORS = {RED, RED2, GRAD_END, "E29C9F", "212529", "223843", "E7E6E6",
                 "8F8F8F", "FFFFFF", "000000"}
OLD_VIOLETS = {"714B67", "7B435D", "724D67", "7030A0"}   # 7030A0 = violet Office des mots mis en valeur
OLD_FONTS = {"Trebuchet MS", "Tahoma", "Calibri", "Calibri Light", "Century Gothic",
             "Montserrat", "Roboto", "Open Sans", "Lato", "Segoe UI"}
ALLOWED_FONTS = {"Aptos", "Aptos Display", "Aptos Narrow", "+mj-lt", "+mn-lt",
                 "+mj-ea", "+mn-ea", "+mj-cs", "+mn-cs",
                 # polices de symboles tolérées (puces, pictos)
                 "Symbol", "Wingdings", "Wingdings 2", "Wingdings 3", "Webdings"}
FONT_BODY = "Aptos"
FONT_TITLE = "Aptos Display"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
R = "{%s}" % NS["r"]

GRADIENT_XML = (
    f'<a:gradFill xmlns:a="{NS["a"]}" rotWithShape="1"><a:gsLst>'
    f'<a:gs pos="0"><a:srgbClr val="{RED}"/></a:gs>'
    f'<a:gs pos="100000"><a:srgbClr val="{GRAD_END}"/></a:gs>'
    '</a:gsLst><a:lin ang="2700000" scaled="1"/></a:gradFill>'
)


def charte_gradient():
    return etree.fromstring(GRADIENT_XML)


# --------------------------------------------------------------------------
# Ouverture de fichiers (.potx -> .pptx)
# --------------------------------------------------------------------------
_TPL_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PRS_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def as_pptx(path: Path) -> Path:
    """Renvoie un chemin .pptx ouvrable par python-pptx (convertit un .potx en cache)."""
    path = Path(path)
    if path.suffix.lower() == ".pptx":
        return path
    CACHE.mkdir(parents=True, exist_ok=True)
    # même nom possible dans 02 et 03 -> suffixe dérivé du chemin complet
    tag = hashlib.sha1(str(path.resolve()).lower().encode()).hexdigest()[:8]
    out = CACHE / f"{path.stem} [{tag}].pptx"
    if out.exists() and out.stat().st_mtime >= path.stat().st_mtime:
        return out
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TPL_CT.encode(), _PRS_CT.encode())
            zout.writestr(item, data)
    return out


def open_prs(path: Path) -> Presentation:
    return Presentation(str(as_pptx(path)))


# --------------------------------------------------------------------------
# Parcours des formes
# --------------------------------------------------------------------------
def iter_shapes(shapes, depth=0, path=""):
    """Parcourt récursivement les formes (y compris l'intérieur des groupes)."""
    for sh in shapes:
        p = f"{path}/{sh.name}" if path else sh.name
        yield sh, depth, p
        if is_group(sh):
            yield from iter_shapes(sh.shapes, depth + 1, p)


def is_group(sh) -> bool:
    return sh._element.tag == P + "grpSp"


def is_picture(sh) -> bool:
    return sh._element.tag == P + "pic"


def is_textbox(sh) -> bool:
    try:
        return sh.shape_type == 17
    except NotImplementedError:
        return False


def is_hidden(slide) -> bool:
    return slide._element.get("show") == "0"


def set_hidden(slide, hidden: bool = True):
    if hidden:
        slide._element.set("show", "0")
    elif "show" in slide._element.attrib:
        del slide._element.attrib["show"]


def notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    tf = slide.notes_slide.notes_text_frame
    return tf.text if tf is not None else ""


def set_notes(slide, text: str):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace(" ", " ").replace("\x0b", " ")).strip()


def shape_texts(slide):
    """Tous les blocs de texte d'une slide (formes, groupes, tableaux)."""
    out = []
    for sh, _, _ in iter_shapes(slide.shapes):
        if sh.is_placeholder and str(sh.placeholder_format.type).split(" ")[0] in (
                "SLIDE_NUMBER", "DATE", "FOOTER"):
            continue
        if re.search(r"num[ée]ro de diapositive|slide number", sh.name, re.I):
            continue  # numéro de slide saisi en zone de texte libre
        if sh.has_text_frame and norm(sh.text_frame.text):
            for para in sh.text_frame.paragraphs:
                t = norm(para.text)
                if t:
                    out.append(t)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        t = norm(para.text)
                        if t:
                            out.append(t)
    return out


def slide_ref(pdoc: Path) -> str:
    """Chemin relatif lisible d'un document par rapport à la racine du projet."""
    try:
        return str(Path(pdoc).resolve().relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(pdoc)


def rel_or_abs(path) -> str:
    """Chemin écrit dans les plans : relatif à PROJECT s'il y est, sinon absolu (fichier joint…)."""
    path = Path(path).resolve()
    try:
        return path.relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def rel_to_source_root(path: Path) -> Path:
    """Chemin relatif d'un fichier par rapport à son dossier racine (02/03/04/05)."""
    path = Path(path).resolve()
    for base in (SRC_FR, REF_EN, RENDER, CORRECTED):
        try:
            return path.relative_to(base.resolve())
        except ValueError:
            continue
    raise ValueError(f"{path} n'est dans aucun dossier 02/03/04/05")


def work_dir(doc: Path) -> Path:
    """Dossier de travail d'un document : work/<chemin relatif sans extension>/"""
    try:
        rel = rel_to_source_root(doc)
        d = WORK / rel.parent / rel.stem
    except ValueError:  # fichier hors 02/03/04/05 (tests, cache…)
        d = WORK / "_autres" / Path(doc).stem
    d.mkdir(parents=True, exist_ok=True)
    return d


def safe_output(path: Path) -> Path:
    """Refuse d'écraser un livrable existant (règle projet : ne jamais écraser)."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(
            f"Refus d'écraser : {path}\n→ Renommer / archiver l'ancienne version ou utiliser --out.")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def emu_in(v) -> float:
    return round(Emu(v).inches, 2) if v is not None else None


def xml_copy(el):
    return copy.deepcopy(el)
