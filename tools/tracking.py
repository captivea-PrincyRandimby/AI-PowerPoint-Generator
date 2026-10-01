"""Bibliothèque des slides récurrentes (unique_slides / unique_slides_EN) + Excel de suivi.

  python tools/tracking.py list [--en]
        -> slides uniques disponibles (nom, titre, empreinte)
  python tools/tracking.py find "<document .pptx|.potx>" [--en]
        -> slides du document identiques (par contenu) à une slide de la bibliothèque
  python tools/tracking.py extract "<doc VALIDÉ 05 .pptx>" <n°> "<Nom de la slide>" [--en]
        -> unique_slides/<Nom>.pptx (1 slide) + ligne "Source" dans l'Excel
  python tools/tracking.py add "<Nom de la slide>" Réutilisation "<document 04>" <position> [--en]
        -> ajoute une ligne dans l'Excel
  python tools/tracking.py sync "<rendu 04 .pptx>" [--en]
        -> ajoute automatiquement les lignes "Réutilisation" d'après build_log.json (mode reuse)

Règle projet : une slide unique est TOUJOURS extraite depuis la version validée (05).
"""
import argparse
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from common import (CORRECTED, ROOT, TEMPLATE, TRACK_EN, TRACK_FR, UNIQUE_EN, UNIQUE_FR, WORK,
                    norm, open_prs, safe_output, shape_texts, slide_ref)
from pptxops import PartCopier, clone_slide, delete_slide

HEADERS = ["Nom de la slide", "Fichier unique", "Rôle", "Document", "Position", "Date"]


def lib(en):
    return (UNIQUE_EN, TRACK_EN) if en else (UNIQUE_FR, TRACK_FR)


def signature(slide):
    """Empreinte de contenu : textes normalisés (ordre trié) -> identité par contenu."""
    txt = sorted(norm(t).lower() for t in shape_texts(slide))
    return hashlib.sha1("\n".join(txt).encode()).hexdigest()[:16], txt


def title_of(slide):
    t = slide.shapes.title
    return norm(t.text_frame.text) if t is not None and t.has_text_frame else ""


def add_row(track: Path, name, fichier, role, document, position):
    track.parent.mkdir(parents=True, exist_ok=True)
    if track.exists():
        wb = load_workbook(track)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Suivi"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="BA131A")
        for col, w in zip("ABCDEF", (38, 38, 14, 70, 10, 12)):
            ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == name and row[2] == role and row[3] == document and str(row[4]) == str(position):
            print(f"(déjà tracé) {name} · {role} · {document} · {position}")
            return
    ws.append([name, fichier, role, document, position, dt.date.today().isoformat()])
    wb.save(track)
    print(f"+ {track.name} : {name} · {role} · {Path(document).name} · pos {position}")


def cmd_list(en):
    folder, _ = lib(en)
    for f in sorted(folder.glob("*.pptx")) if folder.exists() else []:
        s = open_prs(f).slides[0]
        print(f"{f.stem:45} | {title_of(s)[:50]:50} | {signature(s)[0]}")


def library_index(en):
    folder, _ = lib(en)
    idx = {}
    for f in sorted(folder.glob("*.pptx")) if folder.exists() else []:
        s = open_prs(f).slides[0]
        sig, txt = signature(s)
        idx[sig] = (f, txt, title_of(s))
    return idx


def cmd_find(doc, en):
    idx = library_index(en)
    prs = open_prs(doc)
    found = 0
    for n, s in enumerate(prs.slides, 1):
        sig, txt = signature(s)
        if sig in idx:
            print(f"slide {n:3} == {idx[sig][0].stem} (identique)")
            found += 1
            continue
        # correspondance partielle (même titre, >= 70 % des textes communs)
        for f, ltxt, lt in idx.values():
            if lt and lt.lower() == title_of(s).lower():
                common = len(set(txt) & set(ltxt)) / max(len(set(ltxt)), 1)
                if common >= 0.7:
                    print(f"slide {n:3} ~= {f.stem} (titre identique, {common:.0%} des textes)")
                    found += 1
    print(f"{found} correspondance(s)")


def cmd_extract(doc, n, name, en):
    doc = Path(doc).resolve()
    if CORRECTED.resolve() not in doc.parents:
        sys.exit("Refus : une slide unique s'extrait depuis la version VALIDÉE (05. Render corrected).")
    folder, track = lib(en)
    target = safe_output(folder / f"{name}.pptx")
    src = open_prs(doc)
    prs = open_prs(TEMPLATE)
    tpl = list(prs.slides)
    clone_slide(src.slides[n - 1], prs, PartCopier(prs.part.package))
    for s in tpl:
        delete_slide(prs, s)
    prs.save(str(target))
    print(f"-> {target}")
    add_row(track, name, slide_ref(target), "Source", slide_ref(doc), n)


def cmd_sync(render, en):
    render = Path(render).resolve()
    _, track = lib(en)
    for p in WORK.rglob("build_log.json"):
        log = json.loads(p.read_text(encoding="utf-8"))
        if Path(log["output"]).resolve() != render:
            continue
        for e in log["slides"]:
            if e["mode"] == "reuse" and e.get("from"):
                f = Path(e["from"])
                f = f if f.is_absolute() else ROOT / f
                folder, _ = lib(en)
                if folder.resolve() in f.resolve().parents:
                    add_row(track, f.stem, slide_ref(f), "Réutilisation", slide_ref(render), e["out"])
        return
    sys.exit("build_log.json introuvable")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["list", "find", "extract", "add", "sync"])
    ap.add_argument("args", nargs="*")
    ap.add_argument("--en", action="store_true")
    a = ap.parse_args()
    if a.cmd == "list":
        cmd_list(a.en)
    elif a.cmd == "find":
        cmd_find(Path(a.args[0]), a.en)
    elif a.cmd == "extract":
        cmd_extract(a.args[0], int(a.args[1]), a.args[2], a.en)
    elif a.cmd == "add":
        name, role, doc, pos = a.args
        folder, track = lib(a.en)
        add_row(track, name, slide_ref(folder / f"{name}.pptx"), role, slide_ref(Path(doc)), pos)
    else:
        cmd_sync(a.args[0], a.en)


if __name__ == "__main__":
    main()
